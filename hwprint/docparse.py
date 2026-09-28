"""从 Word 原稿里取出首页表头各栏和正文段落。

会议记录类原稿的固定写法：
  第1段  会议名称
  时间：…⇥地点：…
  应到：…⇥实到：…
  主持人：…⇥记录人：…
  会议议题：（后面可能直接跟议题）
  其余（列席人员、授课人、党课主题、活动主题、议题清单、发言）都是正文。
政治理论学习计划原稿：第1段是标题（不写），其余是正文；时间地点人数等每月在界面里填。
"""

import os
import re
from dataclasses import dataclass, field

_SEP = r'(?:\t+|\s{2,}|　+)'
_PAIR_RULES = [
    ('time', 'place', re.compile(r'^\s*时\s*间\s*[：:]\s*(.*?)\s*' + _SEP + r'\s*地\s*点\s*[：:]\s*(.*?)\s*$')),
    ('expected', 'actual', re.compile(r'^\s*应\s*到[^：:\t]*[：:]\s*(.*?)\s*' + _SEP + r'\s*实\s*到[^：:\t]*[：:]\s*(.*?)\s*$')),
    ('host', 'recorder', re.compile(r'^\s*主\s*持\s*人\s*[：:]\s*(.*?)\s*' + _SEP + r'\s*记\s*录\s*人\s*[：:]\s*(.*?)\s*$')),
]
_SINGLE_RULES = [
    ('time', re.compile(r'^\s*时\s*间\s*[：:]\s*(.*?)\s*$')),
    ('place', re.compile(r'^\s*地\s*点\s*[：:]\s*(.*?)\s*$')),
    ('expected', re.compile(r'^\s*应\s*到[^：:]*[：:]\s*(.*?)\s*$')),
    ('actual', re.compile(r'^\s*实\s*到[^：:]*[：:]\s*(.*?)\s*$')),
    ('host', re.compile(r'^\s*主\s*持\s*人\s*[：:]\s*(.*?)\s*$')),
    ('recorder', re.compile(r'^\s*记\s*录\s*人\s*[：:]\s*(.*?)\s*$')),
]
_TOPIC = re.compile(r'^\s*会\s*议\s*议\s*题\s*[：:]\s*(.*?)\s*$')


@dataclass
class Record:
    kind: str                       # 'meeting' / 'study'
    fields: dict = field(default_factory=dict)
    topic: str = ''                 # "会议议题："后面同一行的内容
    body: list = field(default_factory=list)
    source: str = ''
    title: str = ''


def read_docx_paragraphs(path):
    import docx
    d = docx.Document(path)
    out = []
    for p in d.paragraphs:
        t = p.text.replace('\r', '').replace('\xa0', ' ')
        out.extend(t.split('\n'))
    # 原稿偶尔有表格（少见），逐行把单元格文字接上
    for tb in d.tables:
        for r in tb.rows:
            cells = []
            for c in r.cells:
                s = ' '.join(x.strip() for x in c.text.splitlines() if x.strip())
                if s and (not cells or cells[-1] != s):
                    cells.append(s)
            if cells:
                out.append('  '.join(cells))
    return out


def read_paragraphs(path):
    ext = os.path.splitext(path)[1].lower()
    if ext == '.docx':
        return read_docx_paragraphs(path)
    if ext == '.txt':
        raw = open(path, 'rb').read()
        for enc in ('utf-8-sig', 'gb18030'):
            try:
                return raw.decode(enc).splitlines()
            except UnicodeDecodeError:
                continue
    raise ValueError('只支持 .docx 或 .txt 文件（.doc 请先在 Word 里另存为 .docx）')


def guess_kind(path, paras):
    name = os.path.basename(path)
    head = ''.join(paras[:3])
    if '政治理论学习' in name or '政治理论学习' in head:
        return 'study'
    return 'meeting'


def parse(path, kind=None):
    paras = read_paragraphs(path)
    kind = kind or guess_kind(path, paras)
    rec = Record(kind=kind, source=path)
    lines = [p.strip() for p in paras]
    # 去掉开头空行，第一段是标题
    while lines and not lines[0]:
        lines.pop(0)
    if not lines:
        return rec
    rec.title = lines[0]
    rest = lines[1:]
    if kind == 'study':
        rec.body = [p for p in rest if p]
        return rec

    rec.fields['name'] = rec.title
    body = []
    header_done = False
    for idx, line in enumerate(rest):
        if not line:
            continue
        if not header_done:
            hit = False
            for k1, k2, rx in _PAIR_RULES:
                m = rx.match(line)
                if m:
                    rec.fields[k1], rec.fields[k2] = m.group(1), m.group(2)
                    hit = True
                    break
            if not hit:
                for k, rx in _SINGLE_RULES:
                    m = rx.match(line)
                    if m and k not in rec.fields:
                        rec.fields[k] = m.group(1)
                        hit = True
                        break
            if hit:
                continue
            m = _TOPIC.match(line)
            if m:
                rec.topic = m.group(1)
                header_done = True
                continue
            # 表头区的其他行（列席人员、授课人、活动主题等）照原样进正文
            if len(body) >= 4 or idx > 8:
                header_done = True
        body.append(line)
    rec.body = body
    return rec
